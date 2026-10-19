from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "软著申请材料"
OUTPUT_DIR.mkdir(exist_ok=True)
OUTPUT = OUTPUT_DIR / "英语审辨写作训练系统_V1.0_源程序鉴别材料.docx"

SOURCE_FILES = [
    "index.tsx",
    "App.tsx",
    "types.ts",
    "components/LoginScreen.tsx",
    "components/InputSection.tsx",
    "components/PhaseOneCards.tsx",
    "components/ThinkingProcessView.tsx",
    "components/ResultsDisplay.tsx",
    "components/GradingReport.tsx",
    "components/DrillCard.tsx",
    "modules/SocraticCoach.tsx",
    "modules/EssayGrader.tsx",
    "modules/SentenceDrills.tsx",
    "modules/ProfileCenter.tsx",
    "modules/TeacherDashboard.tsx",
    "modules/teacher/OverviewTab.tsx",
    "modules/teacher/AnalyticsTab.tsx",
    "modules/teacher/EssayGalleryTab.tsx",
    "modules/teacher/RosterTab.tsx",
    "modules/teacher/StudentProfilesTab.tsx",
    "modules/teacher/UsageTab.tsx",
    "services/geminiService.ts",
    "services/supabaseDataService.ts",
    "services/authService.ts",
    "services/storageService.ts",
    "services/milestones.ts",
]

PAGES = 60
LINES_PER_PAGE = 60


def set_run_font(run, name="Arial", size=9.0, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    set_run_font(run, size=8)


def build_source_lines():
    lines = []
    counts = []
    for relative_path in SOURCE_FILES:
        path = ROOT / relative_path
        if not path.exists():
            raise FileNotFoundError(path)
        source_lines = [line.rstrip() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
        counts.append((relative_path, len(source_lines)))
        lines.append(f"// ===== SOURCE FILE: {relative_path.replace('/', chr(92))} =====")
        lines.extend(source_lines)
    return lines, counts


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.30)
    section.bottom_margin = Inches(0.30)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)
    section.header_distance = Inches(0.12)
    section.footer_distance = Inches(0.12)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9)

    code_style = doc.styles.add_style("SourceCode", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Arial"
    code_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    code_style.font.size = Pt(9)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = Pt(9.5)
    code_style.paragraph_format.widow_control = False
    code_style.paragraph_format.keep_together = False

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("英语审辨写作训练系统 V1.0")
    set_run_font(header_run, size=8, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("源程序鉴别材料  ")
    set_run_font(footer_run, size=8)
    add_page_field(footer)


def main():
    all_lines, counts = build_source_lines()
    required_lines = PAGES * LINES_PER_PAGE
    half = required_lines // 2
    if len(all_lines) < required_lines:
        raise ValueError(f"候选代码不足 {required_lines} 行，当前仅有 {len(all_lines)} 行")

    selected_lines = all_lines[:half] + all_lines[-half:]
    doc = Document()
    configure_document(doc)

    for index, line in enumerate(selected_lines):
        paragraph = doc.add_paragraph(style="SourceCode")
        paragraph.paragraph_format.left_indent = Inches(0)
        paragraph.paragraph_format.first_line_indent = Inches(0)
        run = paragraph.add_run(line)
        set_run_font(run, size=9)
        if (index + 1) % LINES_PER_PAGE == 0 and index + 1 < len(selected_lines):
            paragraph.add_run().add_break(WD_BREAK.PAGE)

    doc.core_properties.title = "英语审辨写作训练系统 V1.0 源程序鉴别材料"
    doc.core_properties.subject = "软件著作权登记"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    doc.save(OUTPUT)

    print(f"WROTE {OUTPUT}")
    print(f"CANDIDATE_NONEMPTY_LINES {len(all_lines)}")
    for relative_path, count in counts:
        print(f"{count:5d} {relative_path}")


if __name__ == "__main__":
    main()
