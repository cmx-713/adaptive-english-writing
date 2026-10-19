from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "软著申请材料"
OUTPUT_DIR.mkdir(exist_ok=True)
OUTPUT = OUTPUT_DIR / "英语审辨写作训练系统_V1.0_软件使用说明书.docx"


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color="000000"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9E0E8", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(2)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "1F4E79")
        set_cell_borders(cell)
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            if row_index % 2 == 1:
                shade_cell(cell, "F4F7FA")
            set_cell_borders(cell)
            set_cell_text(cell, str(value))
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_screenshot_note(doc, text):
    paragraph = doc.add_paragraph(style="ScreenshotNote")
    run = paragraph.add_run(f"截图提示：{text}")
    set_run_font(run, size=9.5, color="7A4E00")


def add_body(doc, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Inches(0.28)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.left_indent = Inches(0.28)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True, color="000000")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    set_run_font(run, size=9)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in (("Title", 22, "000000"), ("Heading 1", 15, "000000"), ("Heading 2", 12, "000000")):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    note_style = doc.styles.add_style("ScreenshotNote", WD_STYLE_TYPE.PARAGRAPH)
    note_style.font.name = "Microsoft YaHei"
    note_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    note_style.font.size = Pt(9.5)
    note_style.paragraph_format.left_indent = Inches(0.22)
    note_style.paragraph_format.space_before = Pt(4)
    note_style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("英语审辨写作训练系统 V1.0")
    set_run_font(run, size=8.5, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("软件使用说明书  第 ")
    set_run_font(run, size=9)
    add_page_number(footer)
    run = footer.add_run(" 页")
    set_run_font(run, size=9)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(75)
    title.paragraph_format.space_after = Pt(20)
    run = title.add_run("英语审辨写作训练系统")
    set_run_font(run, size=22, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("V1.0 软件使用说明书")
    set_run_font(run, size=14, bold=True, color="1F4E79")

    for _ in range(4):
        doc.add_paragraph()
    cover_info = ["软件形态：具有可视化操作界面的 Web 应用软件", "适用对象：学生、教师及教学管理人员", "文档用途：说明软件的实际操作流程和功能使用方法"]
    for value in cover_info:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, size=11)

    doc.add_page_break()

    add_heading(doc, "目录", 1)
    toc_items = [
        "一 软件概述与运行环境",
        "二 系统入口与身份登录",
        "三 学生工作台与思维训练",
        "四 作文批改与审辨信度分析",
        "五 句子特训与学习反馈",
        "六 学习中心与个人语料库",
        "七 教师工作台与班级管理",
        "八 数据分析与导出",
        "九 系统配置与常见问题",
        "十 功能与代码对应关系",
    ]
    for item in toc_items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        run = paragraph.add_run(item)
        set_run_font(run, size=11)

    doc.add_page_break()

    add_heading(doc, "一 软件概述与运行环境", 1)
    add_heading(doc, "软件用途", 2)
    add_body(doc, "英语审辨写作训练系统面向大学英语写作学习与教学场景，提供从写作话题思考、观点组织、语言支架、作文批改到专项训练和学习分析的连续训练流程。学生可以在系统中形成写作思路并完成作文，教师可以查看班级学习数据、学生作文和审辨信度分析结果。")
    add_heading(doc, "系统组成", 2)
    add_table(doc, ["使用端", "主要功能"], [
        ("学生端", "思维训练、作文批改、审辨信度分析、句子特训、学习中心和个人数据导出"),
        ("教师端", "教学概览、学情分析、学生档案、作文详情、使用统计、名单管理和外校用户管理"),
        ("模型配置", "选择 AI 服务提供商、模型名称、接口地址并保存 API 配置"),
    ], widths=[1.3, 5.7])
    add_heading(doc, "运行环境", 2)
    add_bullet(doc, "客户端使用 Chrome、Edge 或 Firefox 等现代浏览器访问。")
    add_bullet(doc, "系统需要网络连接，以访问数据服务和配置的 AI 模型接口。")
    add_bullet(doc, "本地开发环境使用 Node.js 和 npm；生产环境可部署到支持静态 Web 应用的托管平台。")
    add_bullet(doc, "学生和教师的可见功能由登录角色决定，学生进入学生端，教师进入教师工作台。")
    add_screenshot_note(doc, "系统登录页或系统入口页，展示软件名称、学生登录入口和教师登录入口。")

    add_heading(doc, "二 系统入口与身份登录", 1)
    add_heading(doc, "选择登录角色", 2)
    add_body(doc, "打开系统后进入登录页面。登录页面提供本校学生、教师和外校学生等入口。用户根据身份选择对应标签，填写页面要求的信息后提交。")
    add_heading(doc, "学生登录", 2)
    add_numbered(doc, "选择本校学生或外校学生入口。")
    add_numbered(doc, "选择所属班级或课程。")
    add_numbered(doc, "填写姓名和学号，提交登录。")
    add_numbered(doc, "登录成功后进入学生端思维训练页面，顶部显示当前用户信息。")
    add_body(doc, "学生账号是否可以登录，取决于教师端名单和数据服务中的身份信息。登录失败时，检查班级、姓名和学号是否与名单一致，或联系教师维护名单。")
    add_heading(doc, "教师登录", 2)
    add_numbered(doc, "选择教师入口。")
    add_numbered(doc, "填写教师姓名和教师工号或系统配置的教师账号信息。")
    add_numbered(doc, "提交后进入教师工作台。")
    add_screenshot_note(doc, "登录页面，建议截取完整登录区域，并确保截图中的软件名称与申请名称一致。")

    add_heading(doc, "三 学生工作台与思维训练", 1)
    add_heading(doc, "学生工作台", 2)
    add_body(doc, "学生登录后进入工作台。顶部导航提供思维训练、作文批改、句子特训和学习中心入口。设置按钮用于配置 AI 模型服务，退出按钮用于结束当前本地会话。")
    add_screenshot_note(doc, "学生工作台主界面，显示顶部导航、用户信息、四个功能入口和设置按钮。")
    add_heading(doc, "输入写作话题", 2)
    add_numbered(doc, "进入“思维训练”页面。")
    add_numbered(doc, "在话题输入区域填写需要练习的英语写作题目。")
    add_numbered(doc, "点击开始训练，等待系统生成多个思考角度。")
    add_heading(doc, "选择灵感卡片", 2)
    add_body(doc, "系统根据话题生成多张具有不同维度的灵感卡片。卡片包含论证角度、引导问题、提示信息和可进一步展开的思考方向。学生选择一个卡片后进入该维度的观点训练。")
    add_screenshot_note(doc, "思维训练页面中的话题输入区域和灵感卡片，显示不同论证维度及选择操作。")
    add_heading(doc, "提交观点并获得引导", 2)
    add_numbered(doc, "在观点输入区域用中文或英文写下对当前维度的初步看法。")
    add_numbered(doc, "点击提交观点，系统对观点进行检查并给出反馈。")
    add_numbered(doc, "根据问题提示和论据延伸方向补充观点。")
    add_numbered(doc, "从词汇、搭配和句型支架中选择适合当前观点的表达。")
    add_heading(doc, "完成段落并组合成文", 2)
    add_body(doc, "学生根据句型框架完成英语段落。完成一个维度后可以继续选择其他维度。系统允许在完成至少两个维度后组合成文，将已形成的多个段落整合为完整作文，并保留思维训练过程数据。")
    add_screenshot_note(doc, "观点反馈、词汇搭配支架和句型输入区域；另配一张“组合成文”按钮或完成提示截图。")
    add_heading(doc, "维度宝库", 2)
    add_body(doc, "在思维训练结果区域可以打开“维度宝库”。学生选择经济、环境、社会、科技、健康、心理、教育或伦理等维度后，系统会结合当前话题生成相关联想词，并显示英文表达及中文释义，帮助学生继续拓展论证内容。")
    add_screenshot_note(doc, "思维训练结果页下方的“维度宝库”，显示维度卡片和展开后的联想词。")
    add_heading(doc, "保存记录与历史回顾", 2)
    add_body(doc, "思维训练和作文批改过程会在浏览器本地保存为历史记录。用户可以打开历史列表，按日期和话题查看已保存的思维草稿、观点卡片或作文批改结果；点击记录可以恢复查看，点击删除可以移除不再需要的记录。")
    add_screenshot_note(doc, "思维训练或作文批改页面的历史记录弹窗，显示保存记录、日期、主题、分数或草稿摘要，以及查看和删除操作。")

    add_heading(doc, "四 作文批改与审辨信度分析", 1)
    add_heading(doc, "提交作文批改", 2)
    add_numbered(doc, "点击顶部导航中的“作文批改”。")
    add_numbered(doc, "在输入框中输入或粘贴英语作文，也可以从思维训练结果跳转并自动带入作文。")
    add_numbered(doc, "点击开始批改，等待系统完成分析。")
    add_numbered(doc, "在批改报告中查看评分、问题诊断、修改建议和参考表达。")
    add_heading(doc, "查看批改报告", 2)
    add_table(doc, ["报告区域", "内容"], [
        ("评分概览", "总分、分项得分、总体评价和问题分级"),
        ("深度精批", "逐段或逐句的问题说明、修改建议和语言分析"),
        ("范文对比", "学生表达与参考表达的对照，以及表达提升说明"),
        ("学练一体", "根据批改结果进入语法、句子或结构专项训练"),
    ], widths=[1.5, 5.5])
    add_screenshot_note(doc, "作文批改报告页，建议包含评分概览、问题列表和报告标签页。")
    add_heading(doc, "审辨信度分析", 2)
    add_body(doc, "完成思维训练并组合成文后，系统可以根据学生的完整思维过程进行审辨信度分析。分析结果包括观点一致性、论证递进性、语言自主性和观点拓展度等维度，并展示综合评分和诊断说明。")
    add_body(doc, "学生可以展开各维度查看评分依据。对于薄弱维度，系统提供进入句子特训的入口，形成“分析—训练—再练习”的反馈路径。")
    add_screenshot_note(doc, "审辨信度分析页面，显示四个评分维度、综合结果、评分说明和进入专项训练的按钮。")

    add_heading(doc, "五 句子特训与学习反馈", 1)
    add_heading(doc, "进入句子特训", 2)
    add_body(doc, "句子特训可以从顶部导航直接进入，也可以从作文批改报告或审辨信度分析结果中的专项训练入口进入。由报告跳转时，系统会结合当前错误或薄弱环节准备训练内容。")
    add_heading(doc, "选择训练模式", 2)
    add_table(doc, ["训练模式", "训练目的"], [
        ("语法修正", "识别并改正句子中的语法问题"),
        ("句子提升", "将一般表达提升为更准确、更正式的学术表达"),
        ("结构训练", "练习复杂句型、连接关系和论证结构"),
    ], widths=[1.5, 5.5])
    add_heading(doc, "完成专项练习", 2)
    add_numbered(doc, "选择训练模式并查看系统提供的练习句。")
    add_numbered(doc, "根据题目要求输入修改后的句子或选择合适表达。")
    add_numbered(doc, "提交答案，查看系统反馈和解释。")
    add_numbered(doc, "根据反馈继续下一道练习，训练结果会保存到学习数据中。")
    add_screenshot_note(doc, "句子特训页面，显示三种训练模式和一道具体练习题。")

    add_heading(doc, "六 学习中心与个人语料库", 1)
    add_heading(doc, "查看学习积累", 2)
    add_body(doc, "学习中心集中展示学生在训练过程中积累的词汇和地道搭配，并显示思维训练、作文批改和句子特训次数。学生可以按话题筛选内容，也可以通过搜索快速查找已保存的表达。词汇银行中的词条支持删除，删除后不会影响已经完成的作文或训练记录。")
    add_heading(doc, "查看成长分析", 2)
    add_body(doc, "学习中心展示最近作文的分数走势、能力雷达、四个评分维度的历史趋势、首次与最近成绩对比，以及基于薄弱维度生成的学习建议。完成相应条件后，学习里程碑徽章会自动点亮。审辨信度记录还会展示单次评分卡、历史趋势、总体评语和四维得分明细。")
    add_screenshot_note(doc, "学习中心的统计卡片、学习里程碑、写作分数走势、能力雷达、成长对比和审辨信度趋势区域。")
    add_heading(doc, "导出个人语料库", 2)
    add_numbered(doc, "进入学习中心。")
    add_numbered(doc, "查看词汇银行或地道搭配。")
    add_numbered(doc, "使用导出功能生成 CSV 或 Markdown 文件。")
    add_numbered(doc, "将导出的内容用于个人复习、写作积累或后续学习整理。")
    add_screenshot_note(doc, "学习中心页面，显示词汇银行、地道搭配、筛选搜索区域和导出按钮。")

    add_heading(doc, "七 教师工作台与班级管理", 1)
    add_heading(doc, "教师工作台", 2)
    add_body(doc, "教师登录后进入教师工作台。页面顶部提供教学概览、学情分析、学生档案、作文详情、使用统计、受邀名单和外校用户等功能入口。工作台支持按班级筛选数据。")
    add_screenshot_note(doc, "教师工作台首页，显示顶部功能标签、班级筛选、学生人数、作文数量、特训次数和操作按钮。")
    add_heading(doc, "班级和名单管理", 2)
    add_numbered(doc, "进入“受邀名单”。")
    add_numbered(doc, "填写班级名称，选择本校学生或外校学生名单类型。")
    add_numbered(doc, "按页面提示批量粘贴学生信息。")
    add_numbered(doc, "提交导入或更新名单。")
    add_numbered(doc, "在当前名单中按班级筛选并查看已导入学生。")
    add_body(doc, "名单信息用于学生登录匹配。班级名称应与学生登录时选择的班级保持一致。")
    add_screenshot_note(doc, "受邀名单页面，显示新建班级、名单类型、批量导入区域和当前名单。")
    add_heading(doc, "学生档案与作文详情", 2)
    add_body(doc, "学生档案页面按班级展示学生信息和学习次数，支持查看个人学习表现、作文批改次数、句子特训次数、思维训练次数及审辨信度结果。教师还可以在学生档案中为未分班学生分配班级。作文详情页面用于浏览学生已保存的作文分析记录。")
    add_screenshot_note(doc, "学生档案或作文详情页面，显示班级筛选、学生列表和选中学生的学习数据。")
    add_heading(doc, "外校用户管理", 2)
    add_body(doc, "教师工作台提供外校用户查看入口，用于查看已注册的外校用户。外校用户仅使用学生端功能，不能访问教师工作台。")

    add_heading(doc, "八 数据分析与导出", 1)
    add_heading(doc, "教学概览", 2)
    add_body(doc, "教学概览展示班级人数、作文数量、句子特训次数等基本统计，并提供班级横向对比和班级整体进步分析。教师可以通过班级选择器查看全部班级或某一班级的数据。")
    add_heading(doc, "学情分析", 2)
    add_body(doc, "学情分析从四个能力维度比较班级平均水平与评分标准，展示多次批改学生的进步追踪，并统计各维度低于班级平均水平的学生人数。")
    add_heading(doc, "批量审辨信度分析", 2)
    add_numbered(doc, "点击教师工作台中的“批量审辨信度”。")
    add_numbered(doc, "选择目标班级。")
    add_numbered(doc, "选择是否包含已有审辨信度结果。")
    add_numbered(doc, "开始批量分析，完成后返回学情分析页面查看结果。")
    add_heading(doc, "复核审辨信度结果", 2)
    add_body(doc, "教师查看学生思维过程和审辨信度结果时，可以对自动生成的评分进行重新分析，也可以点击“标记已复核”记录教师已经检查过的结果。复核标记会在教师端保留，便于区分自动分析结果和教师确认结果。")
    add_heading(doc, "数据导出", 2)
    add_body(doc, "教师可以导出学生学习汇总和审辨信度汇总。导出数据会根据当前班级筛选范围生成 CSV 文件，用于教学记录整理和班级分析。")
    add_screenshot_note(doc, "教师工作台中的批量审辨信度弹窗和导出数据菜单。")

    add_heading(doc, "九 系统配置与常见问题", 1)
    add_heading(doc, "配置 AI 模型服务", 2)
    add_body(doc, "点击学生端顶部的设置按钮，打开 API Settings。系统支持 DeepSeek、通义千问、智谱 GLM 和自定义 OpenAI 兼容接口。用户可以配置模型名称、API Base URL 和 API Key，并点击 Save Configuration 保存。配置内容保存在当前浏览器本地。")
    add_screenshot_note(doc, "API Settings 设置窗口，显示模型提供商、模型名称、API Base URL、API Key 和保存按钮。截图中的密钥必须打码或使用空值。")
    add_heading(doc, "常见问题", 2)
    add_table(doc, ["问题", "处理方法"], [
        ("无法登录", "确认班级、姓名、学号或教师账号与系统名单一致；必要时联系教师维护名单。"),
        ("AI 功能无响应", "检查网络、API Key、接口地址和模型名称，确认模型服务可用。"),
        ("批改或分析失败", "检查输入是否完整，稍后重试；仍失败时检查浏览器控制台和数据服务状态。"),
        ("学习中心为空", "先完成思维训练或其他会产生学习积累的功能，再返回学习中心查看。"),
        ("导出文件乱码", "用 Excel 的“从文本/CSV 导入”功能打开，并选择 UTF-8 编码。"),
        ("教师端看不到数据", "检查班级筛选条件和数据服务连接，确认学生已完成并保存相关训练。"),
    ], widths=[1.7, 5.3])

    add_heading(doc, "十 功能与代码对应关系", 1)
    add_body(doc, "本说明书按当前项目源代码和实际页面组织。下表用于核对说明书中的核心功能与实现模块，正式使用时不需要用户阅读源代码。")
    add_table(doc, ["功能", "主要实现模块"], [
        ("应用入口、角色分流和页面导航", "index.tsx、App.tsx"),
        ("学生、教师和外校用户登录", "components/LoginScreen.tsx、services/authService.ts、services/supabaseDataService.ts"),
        ("思维训练和灵感卡片", "modules/SocraticCoach.tsx、components/InputSection.tsx、components/PhaseOneCards.tsx"),
        ("观点反馈、思维过程和维度宝库", "components/ResultsDisplay.tsx、components/ThinkingProcessView.tsx、components/DimensionDictionary.tsx"),
        ("思维训练和作文历史记录", "components/HistoryModal.tsx、services/storageService.ts"),
        ("作文批改和报告展示", "modules/EssayGrader.tsx、components/GradingReport.tsx"),
        ("句子特训", "modules/SentenceDrills.tsx、components/DrillCard.tsx"),
        ("学习中心、历史趋势、里程碑、词汇搭配和导出", "modules/ProfileCenter.tsx、services/storageService.ts、services/milestones.ts"),
        ("教师工作台和班级管理", "modules/TeacherDashboard.tsx、modules/teacher/RosterTab.tsx"),
        ("教学概览、学情分析、学生分班和教师复核", "modules/teacher/OverviewTab.tsx、AnalyticsTab.tsx、StudentProfilesTab.tsx、components/ThinkingProcessView.tsx"),
        ("作文详情和使用统计", "modules/teacher/EssayGalleryTab.tsx、UsageTab.tsx"),
        ("AI 调用、结构化结果和审辨信度分析", "services/geminiService.ts"),
        ("数据保存、查询和统计", "services/supabaseDataService.ts、services/supabaseClient.ts"),
    ], widths=[2.4, 4.6])

    add_body(doc, "本说明书只描述当前版本中能够通过页面操作找到并运行的功能。接口服务、数据库和模型配置属于系统运行所需的支撑条件，具体使用效果取决于网络连接、服务配置和权限状态。")

    doc.core_properties.title = "英语审辨写作训练系统 V1.0 软件使用说明书"
    doc.core_properties.subject = "软件著作权登记文档"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    doc.save(OUTPUT)
    print(f"WROTE {OUTPUT}")


if __name__ == "__main__":
    build_document()
